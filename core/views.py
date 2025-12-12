from weasyprint import HTML
from django.db import transaction
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Count, Q
from django.core.paginator import Paginator
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from datetime import timedelta, datetime
import json
import logging
from django.template.loader import render_to_string
from core.models import Soldier, Leave, LeaveType, ServiceLocation, ServiceStatus, SystemSetting, HealthStatus
from core.forms import SettingsForm
from jdatetime import datetime as jdatetime

logger = logging.getLogger('core')


def parse_date_input(value):
    """
    Accepts Jalali (YYYY/MM/DD) or Gregorian (YYYY-MM-DD) strings and returns (gregorian_date, jalali_display).
    """
    value = (value or '').strip()
    if not value:
        return None, ''
    try:
        if '/' in value:
            jd = jdatetime.strptime(value, '%Y/%m/%d')
            greg = jd.togregorian().date()
        else:
            greg = datetime.fromisoformat(value).date()
        jalali_display = jdatetime.fromgregorian(date=greg).strftime('%Y/%m/%d')
        return greg, jalali_display
    except Exception:
        return None, value


def panel_login(request):
    if request.user.is_authenticated:
        return redirect('panel_dashboard')

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')

        if not username or not password:
            messages.error(request, 'لطفاً نام کاربری و رمز عبور را وارد کنید.')
            return render(request, 'dashboard/login.html')

        user = authenticate(request, username=username, password=password)
        if user is not None and user.is_staff:
            login(request, user)
            logger.info(f'User {username} logged in successfully')
            return redirect('panel_dashboard')
        else:
            logger.warning(f'Failed login attempt for username: {username}')
            messages.error(request, 'نام کاربری یا رمز عبور اشتباه است یا دسترسی ندارید.')

    return render(request, 'dashboard/login.html')


@login_required
def panel_logout(request):
    username = request.user.username
    logout(request)
    logger.info(f'User {username} logged out')
    return redirect('panel_login')


@login_required
def panel_dashboard(request):
    today = timezone.now().date()
    ninety_days_later = today + timedelta(days=90)

    # محاسبه آماری که نیاز به service_end_date دارن، در پایتون انجام می‌دیم
    # Optimized: Load all soldiers once with prefetch_related to reduce queries
    soldiers = list(Soldier.objects.select_related('service_status').prefetch_related('reductions').all())
    total_soldiers = len(soldiers)

    on_leave_today = Leave.objects.filter(
        start_date__lte=today,
        end_date__gte=today
    ).count()

    # محاسبه سربازانی که تا ۹۰ روز آینده خدمت‌شون تموم می‌شه
    ending_soon = 0
    for soldier in soldiers:
        try:
            end_date = soldier.service_end_date()
            if end_date and today <= end_date <= ninety_days_later:
                ending_soon += 1
        except Exception as e:
            logger.warning(f'Error calculating service_end_date for soldier {soldier.id}: {e}')
            continue

    # Chart data: Leave type distribution
    leave_stats = Leave.objects.values('leave_type').annotate(
        total_days=Count('days')
    ).order_by('leave_type')

    leave_chart_data = {
        'labels': [],
        'data': [],
        'colors': ['#dc2626', '#16a34a', '#3b82f6']
    }

    for stat in leave_stats:
        leave_type = stat['leave_type']
        leave_chart_data['labels'].append(dict(LeaveType.choices).get(leave_type, leave_type))
        leave_chart_data['data'].append(stat['total_days'])

    # Chart data: Service ending in next 6 months
    six_months_later = today + timedelta(days=180)
    ending_chart_data = {
        'labels': [],
        'data': []
    }

    # Group by month
    current_month = today
    month_labels = []
    month_counts = []

    for i in range(6):
        month_start = current_month.replace(day=1)
        if i < 5:
            if month_start.month == 12:
                month_end = month_start.replace(year=month_start.year + 1, month=1)
            else:
                month_end = month_start.replace(month=month_start.month + 1)
        else:
            month_end = six_months_later

        # Count soldiers ending in this month
        # Use prefetched soldiers list to avoid multiple queries
        count = 0
        for soldier in soldiers:
            try:
                end_date = soldier.service_end_date()
                if end_date and month_start <= end_date < month_end:
                    count += 1
            except:
                continue

        # Convert to Jalali month name
        try:
            jd = jdatetime.fromgregorian(date=month_start)
            month_labels.append(jd.strftime('%Y/%m'))
        except:
            month_labels.append(month_start.strftime('%Y/%m'))

        month_counts.append(count)
        current_month = month_end

    ending_chart_data['labels'] = month_labels
    ending_chart_data['data'] = month_counts

    stats = {
        'total_soldiers': total_soldiers,
        'on_leave_today': on_leave_today,
        'ending_soon': ending_soon,
    }

    context = {
        'stats': stats,
        'leave_chart_data': json.dumps(leave_chart_data),
        'ending_chart_data': json.dumps(ending_chart_data),
    }

    return render(request, 'dashboard/dashboard.html', context)


@login_required
def panel_soldiers(request):
    show_completed = request.GET.get('show_completed') == '1'
    sort_key = request.GET.get('sort', '-start_date')
    page_size_raw = request.GET.get('page_size')

    try:
        page_size = int(page_size_raw)
        page_size = max(5, min(page_size, 100))
    except (TypeError, ValueError):
        page_size = 25

    sort_map = {
        'name': ('last_name', 'first_name'),
        '-name': ('-last_name', '-first_name'),
        'national_code': 'national_code',
        '-national_code': '-national_code',
        'start_date': 'service_start_date',
        '-start_date': '-service_start_date',
        'dossier': 'dossier_code',
        '-dossier': '-dossier_code',
        'status': 'service_status__title',
        '-status': '-service_status__title',
        'location': 'current_location__name',
        '-location': '-current_location__name',
        'created': '-created_at',
        '-created': '-created_at',
    }
    order_fields = sort_map.get(sort_key, ('-service_start_date',))
    if isinstance(order_fields, str):
        order_fields = (order_fields,)

    soldiers = Soldier.objects.select_related('service_status', 'current_location')
    if not show_completed:
        soldiers = soldiers.filter(Q(service_status__isnull=True) | Q(service_status__is_active_service=True) | Q(
            service_status__is_active_service__isnull=True))
    soldiers = soldiers.order_by(*order_fields)

    paginator = Paginator(soldiers, page_size)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    query_params = request.GET.copy()
    for key in ('page', 'start_from_display', 'start_to_display'):
        if key in query_params:
            query_params.pop(key)
    if 'page_size' not in query_params:
        query_params['page_size'] = page_size
    base_querystring = query_params.urlencode()

    return render(request, 'dashboard/soldiers_list.html', {
        'soldiers': page_obj,
        'page_obj': page_obj,
        'show_completed': show_completed,
        'page_size': page_size,
        'sort': sort_key if sort_key in sort_map else '-start_date',
        'base_querystring': base_querystring,
    })


@login_required
def panel_leaves(request):
    leaves = Leave.objects.select_related('soldier').order_by('-created_at')

    # Pagination
    paginator = Paginator(leaves, 25)  # Show 25 leaves per page
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'dashboard/leaves_list.html', {
        'leaves': page_obj,
        'page_obj': page_obj,
    })


@login_required
def panel_search(request):
    soldiers = Soldier.objects.select_related('service_status', 'current_location', 'health_status').all()

    # Filters
    query = request.GET.get('q', '').strip()
    dispatch_type = request.GET.get('dispatch_type', '')
    service_status_id = request.GET.get('service_status', '')
    location_id = request.GET.get('location', '')
    health_status_id = request.GET.get('health_status', '')
    start_from_raw = request.GET.get('start_from', '')
    start_to_raw = request.GET.get('start_to', '')
    show_completed = request.GET.get('show_completed') == '1'
    page_size_raw = request.GET.get('page_size')

    try:
        page_size = int(page_size_raw)
        page_size = max(5, min(page_size, 100))
    except (TypeError, ValueError):
        page_size = 25

    start_from_date, start_from_display = parse_date_input(start_from_raw)
    start_to_date, start_to_display = parse_date_input(start_to_raw)

    if not show_completed:
        soldiers = soldiers.filter(Q(service_status__isnull=True) | Q(service_status__is_active_service=True) | Q(
            service_status__is_active_service__isnull=True))

    if query:
        soldiers = soldiers.filter(
            Q(first_name__icontains=query) |
            Q(last_name__icontains=query) |
            Q(national_code__icontains=query) |
            Q(dossier_code__icontains=query) |
            Q(phone_personal__icontains=query) |
            Q(phone_parent__icontains=query) |
            Q(phone_home__icontains=query) |
            Q(phone_work__icontains=query) |
            Q(home_address__icontains=query)
        )

    if dispatch_type:
        soldiers = soldiers.filter(dispatch_type=dispatch_type)

    if service_status_id:
        soldiers = soldiers.filter(service_status_id=service_status_id)

    if location_id:
        soldiers = soldiers.filter(current_location_id=location_id)

    if health_status_id:
        soldiers = soldiers.filter(health_status_id=health_status_id)

    if start_from_date:
        soldiers = soldiers.filter(service_start_date__gte=start_from_date)
    if start_to_date:
        soldiers = soldiers.filter(service_start_date__lte=start_to_date)

    # Pagination with preserved query params
    paginator = Paginator(soldiers, page_size)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    query_params = request.GET.copy()
    if 'page' in query_params:
        query_params.pop('page')
    if 'page_size' not in query_params:
        query_params['page_size'] = page_size
    base_querystring = query_params.urlencode()

    from core.models import DispatchType

    context = {
        'soldiers': page_obj,
        'page_obj': page_obj,
        'query': query,
        'dispatch_types': DispatchType.choices,
        'service_statuses': ServiceStatus.objects.all(),
        'locations': ServiceLocation.objects.all(),
        'health_statuses': HealthStatus.objects.all(),
        'selected_dispatch_type': dispatch_type,
        'selected_service_status': service_status_id,
        'selected_location': location_id,
        'selected_health_status': health_status_id,
        'start_from_display': start_from_display or start_from_raw,
        'start_to_display': start_to_display or start_to_raw,
        'start_from_value': (start_from_date.isoformat() if start_from_date else start_from_raw),
        'start_to_value': (start_to_date.isoformat() if start_to_date else start_to_raw),
        'show_completed': show_completed,
        'base_querystring': base_querystring,
        'page_size': page_size,
    }

    return render(request, 'dashboard/soldier_search.html', context)


def page_not_found(request, exception):
    return render(request, '404.html', status=404)


def server_error(request):
    return render(request, '500.html', status=500)


@login_required
def panel_soldier_detail(request, soldier_id):
    soldier = get_object_or_404(Soldier, id=soldier_id)
    leaves = soldier.leaves.all().order_by('-start_date')
    reductions = soldier.reductions.all()
    balances = soldier.leave_balances()

    context = {
        'soldier': soldier,
        'leaves': leaves,
        'reductions': reductions,
        'leave_balances': balances,
    }

    return render(request, 'dashboard/soldier_detail.html', context)


@login_required
def soldier_profile_pdf(request, soldier_id):
    """Generate PDF profile for a soldier (A5 size)"""
    soldier = get_object_or_404(Soldier, id=soldier_id)
    balances = soldier.leave_balances()
    now = timezone.now()

    html = render_to_string('pdf/soldier_profile.html', {
        'soldier': soldier,
        'balances': balances,
        'request': request,
        'now': now,
    })
    pdf_file = HTML(string=html, base_url=request.build_absolute_uri()).write_pdf()
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="soldier_{soldier.id}_profile.pdf"'
    return response


@login_required
def soldier_leaves_pdf(request, soldier_id):
    """Generate PDF report of all leaves for a soldier (A4)"""
    soldier = get_object_or_404(Soldier, id=soldier_id)
    leaves = soldier.leaves.all().order_by('-start_date')
    balances = soldier.leave_balances()
    now = timezone.now()

    html = render_to_string('pdf/soldier_leaves.html', {
        'soldier': soldier,
        'leaves': leaves,
        'balances': balances,
        'request': request,
        'now': now,
    })
    pdf_file = HTML(string=html, base_url=request.build_absolute_uri()).write_pdf()
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="soldier_{soldier.id}_leaves.pdf"'
    return response


@login_required
def panel_soldier_delete(request, soldier_id):
    soldier = get_object_or_404(Soldier, id=soldier_id)

    if request.method == 'POST':
        soldier_name = soldier.full_name()
        soldier_id = soldier.id
        soldier.delete()
        logger.info(f'Soldier {soldier_id} ({soldier_name}) deleted by user {request.user.username}')
        messages.success(request, f'سرباز "{soldier_name}" با موفقیت حذف شد.')
        return redirect('panel_soldiers')

    return render(request, 'dashboard/soldier_confirm_delete.html', {'soldier': soldier})


@login_required
def panel_leave_delete(request, leave_id):
    leave = get_object_or_404(Leave, id=leave_id)

    if request.method == 'POST':
        soldier_name = leave.soldier.full_name()
        leave_id = leave.id
        leave.delete()
        logger.info(f'Leave {leave_id} for soldier {soldier_name} deleted by user {request.user.username}')
        messages.success(request, f'مرخصی سرباز "{soldier_name}" با موفقیت حذف شد.')
        return redirect('panel_leaves')

    return render(request, 'dashboard/leave_confirm_delete.html', {'leave': leave})


@login_required
def health_check(request):
    """Health check endpoint for monitoring"""
    from django.db import connection
    try:
        # Check database connection
        with connection.cursor() as cursor:
            cursor.execute("SELECT 1")

        return JsonResponse({
            'status': 'healthy',
            'database': 'connected',
            'timestamp': timezone.now().isoformat()
        })
    except Exception as e:
        logger.error(f'Health check failed: {e}')
        return JsonResponse({
            'status': 'unhealthy',
            'database': 'disconnected',
            'error': str(e),
            'timestamp': timezone.now().isoformat()
        }, status=503)


@login_required
def export_soldiers_excel(request):
    """Export soldiers to Excel"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        from django.http import HttpResponse

        soldiers = Soldier.objects.select_related('service_status', 'current_location', 'health_status').all()

        wb = Workbook()
        ws = wb.active
        ws.title = "سربازان"

        # Headers
        headers = ['نام', 'نام خانوادگی', 'کد ملی', 'کد پرونده', 'نوع اعزام',
                   'تاریخ شروع خدمت', 'پایان خدمت', 'وضعیت خدمت', 'محل خدمت']
        ws.append(headers)

        # Style headers
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

        # Data
        for soldier in soldiers:
            end_date = soldier.service_end_date()
            ws.append([
                soldier.first_name,
                soldier.last_name,
                soldier.national_code,
                soldier.dossier_code or '',
                soldier.get_dispatch_type_display(),
                soldier.service_start_date.strftime('%Y/%m/%d') if soldier.service_start_date else '',
                end_date.strftime('%Y/%m/%d') if end_date else '',
                soldier.service_status.title if soldier.service_status else '',
                soldier.current_location.name if soldier.current_location else ''
            ])

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="soldiers.xlsx"'
        wb.save(response)
        logger.info(f'Soldiers exported to Excel by user {request.user.username}')
        return response
    except ImportError:
        messages.error(request, 'برای استفاده از این قابلیت، openpyxl را نصب کنید: pip install openpyxl')
        return redirect('panel_soldiers')
    except Exception as e:
        logger.error(f'Error exporting soldiers: {e}')
        messages.error(request, 'خطا در خروجی گرفتن اطلاعات')
        return redirect('panel_soldiers')


@login_required
def export_leaves_excel(request):
    """Export leaves to Excel"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        from django.http import HttpResponse

        leaves = Leave.objects.select_related('soldier').order_by('-start_date')

        wb = Workbook()
        ws = wb.active
        ws.title = "مرخصی‌ها"

        # Headers
        headers = ['نام سرباز', 'کد ملی', 'نوع مرخصی', 'تاریخ رفت',
                   'تاریخ برگشت', 'مدت (روز)']
        ws.append(headers)

        # Style headers
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

        # Data
        for leave in leaves:
            ws.append([
                leave.soldier.full_name(),
                leave.soldier.national_code,
                leave.get_leave_type_display(),
                leave.start_date.strftime('%Y/%m/%d'),
                leave.end_date.strftime('%Y/%m/%d'),
                leave.days
            ])

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="leaves.xlsx"'
        wb.save(response)
        logger.info(f'Leaves exported to Excel by user {request.user.username}')
        return response
    except ImportError:
        messages.error(request, 'برای استفاده از این قابلیت، openpyxl را نصب کنید: pip install openpyxl')
        return redirect('panel_leaves')
    except Exception as e:
        logger.error(f'Error exporting leaves: {e}')
        messages.error(request, 'خطا در خروجی گرفتن اطلاعات')
        return redirect('panel_leaves')


@login_required
def panel_reports(request):
    """Reports page"""
    from django.db.models import Count, Sum
    from datetime import datetime

    # Monthly leave statistics
    today = timezone.now().date()
    current_month_start = today.replace(day=1)

    # Get last 6 months
    monthly_leaves = []
    for i in range(6):
        month_start = current_month_start
        if month_start.month == 1:
            month_start = month_start.replace(year=month_start.year - 1, month=12)
        else:
            month_start = month_start.replace(month=month_start.month - 1)

        month_end = current_month_start
        if i == 0:
            month_end = today

        leaves_in_month = Leave.objects.filter(
            start_date__gte=month_start,
            start_date__lt=month_end
        )

        monthly_leaves.append({
            'month': month_start.strftime('%Y/%m'),
            'total': leaves_in_month.count(),
            'total_days': leaves_in_month.aggregate(Sum('days'))['days__sum'] or 0,
        })

        current_month_start = month_start

    monthly_leaves.reverse()

    # Service completion forecast
    six_months_later = today + timedelta(days=180)
    ending_forecast = []
    current_month = today.replace(day=1)

    for i in range(6):
        if current_month.month == 12:
            month_end = current_month.replace(year=current_month.year + 1, month=1)
        else:
            month_end = current_month.replace(month=current_month.month + 1)

        if i == 5:
            month_end = six_months_later

        count = 0
        soldiers = Soldier.objects.select_related('service_status').prefetch_related('reductions').all()
        for soldier in soldiers:
            try:
                end_date = soldier.service_end_date()
                if end_date and current_month <= end_date < month_end:
                    count += 1
            except:
                continue

        try:
            jd = jdatetime.fromgregorian(date=current_month)
            month_label = jd.strftime('%Y/%m')
        except:
            month_label = current_month.strftime('%Y/%m')

        ending_forecast.append({
            'month': month_label,
            'count': count
        })

        current_month = month_end

    # Soldier distribution by location
    location_dist = ServiceLocation.objects.annotate(
        soldier_count=Count('soldier')
    ).filter(soldier_count__gt=0).order_by('-soldier_count')[:10]

    # Soldier distribution by status
    status_dist = ServiceStatus.objects.annotate(
        soldier_count=Count('soldier')
    ).order_by('-soldier_count')

    context = {
        'monthly_leaves': monthly_leaves,
        'ending_forecast': ending_forecast,
        'location_dist': location_dist,
        'status_dist': status_dist,
    }

    return render(request, 'dashboard/reports.html', context)


@login_required
def panel_settings(request):
    """System settings page"""

    def _get_bool(key, default=True):
        val = SystemSetting.get_value(key, 'true' if default else 'false')
        return str(val).lower() == 'true'

    initial = {
        'base_service_months': int(SystemSetting.get_value('BASE_SERVICE_MONTHS', '21') or 21),
        'enable_leave_entitlement': _get_bool('ENABLE_LEAVE_ENTITLEMENT', True),
        'enable_leave_medical': _get_bool('ENABLE_LEAVE_MEDICAL', True),
        'enable_leave_reward': _get_bool('ENABLE_LEAVE_REWARD', True),
    }

    if request.method == 'POST':
        form = SettingsForm(request.POST)
        if form.is_valid():
            SystemSetting.objects.update_or_create(key='BASE_SERVICE_MONTHS',
                                                   defaults={'value': str(form.cleaned_data['base_service_months'])})
            SystemSetting.objects.update_or_create(key='ENABLE_LEAVE_ENTITLEMENT', defaults={
                'value': str(form.cleaned_data['enable_leave_entitlement'])})
            SystemSetting.objects.update_or_create(key='ENABLE_LEAVE_MEDICAL',
                                                   defaults={'value': str(form.cleaned_data['enable_leave_medical'])})
            SystemSetting.objects.update_or_create(key='ENABLE_LEAVE_REWARD',
                                                   defaults={'value': str(form.cleaned_data['enable_leave_reward'])})
            messages.success(request, 'تنظیمات با موفقیت ذخیره شد.')
            return redirect('panel_settings')
    else:
        form = SettingsForm(initial=initial)

    return render(request, 'dashboard/settings.html', {'form': form})


@login_required
def seed_data(request):
    """Seed database with fake data for testing

    WARNING: This endpoint should be removed or commented in production!
    """
    from django.utils import timezone
    from datetime import timedelta
    import random
    from core.models import (
        ServiceLocation, HealthStatus, ServiceStatus, Expertise,
        ServiceReductionType, Soldier, Leave, SoldierServiceReduction,
        DispatchType, LeaveType
    )

    if request.method != 'POST':
        return JsonResponse({'error': 'Only POST method allowed'}, status=405)

    try:
        with transaction.atomic():
            locations = [
                ServiceLocation.objects.get_or_create(name='پادگان مرکزی')[0],
                ServiceLocation.objects.get_or_create(name='واحد آموزش')[0],
                ServiceLocation.objects.get_or_create(name='واحد عملیات')[0],
                ServiceLocation.objects.get_or_create(name='واحد پشتیبانی')[0],
            ]

            health_statuses = [
                HealthStatus.objects.get_or_create(
                    title='سالم',
                    defaults={'is_exempt_from_combat': False}
                )[0],
                HealthStatus.objects.get_or_create(
                    title='معاف از رزم',
                    defaults={'is_exempt_from_combat': True}
                )[0],
                HealthStatus.objects.get_or_create(
                    title='محدود',
                    defaults={'is_exempt_from_combat': False}
                )[0],
            ]

            # --- IMPORTANT PART: separate active vs finished statuses ---
            status_active_service = ServiceStatus.objects.get_or_create(
                title='در حال خدمت',
                defaults={'is_active_service': True}
            )[0]
            status_finished_service = ServiceStatus.objects.get_or_create(
                title='پایان خدمت',
                defaults={'is_active_service': False}
            )[0]
            status_on_leave = ServiceStatus.objects.get_or_create(
                title='مرخص',
                defaults={'is_active_service': True}
            )[0]

            # For seeding, only pick from active statuses to avoid dossier_code = -1
            active_statuses_for_seed = [status_active_service, status_on_leave]
            # finished status still exists in DB, but we don't use it in seed

            expertises = [
                Expertise.objects.get_or_create(title='رایانه')[0],
                Expertise.objects.get_or_create(title='برق')[0],
                Expertise.objects.get_or_create(title='مکانیک')[0],
                Expertise.objects.get_or_create(title='رانندگی')[0],
            ]

            reduction_types = [
                ServiceReductionType.objects.get_or_create(
                    title='آموزشی',
                    defaults={'default_days': 30}
                )[0],
                ServiceReductionType.objects.get_or_create(
                    title='خدمات عمومی',
                    defaults={'default_days': 15}
                )[0],
            ]

            first_names = ['علی', 'محمد', 'حسن', 'حسین', 'رضا', 'امیر', 'سعید', 'مهدی', 'احمد', 'حامد',
                           'فرهاد', 'کامران', 'بهرام', 'آرش', 'پویا', 'نیما', 'سینا', 'آرمان', 'امیرحسین', 'محمدحسین']
            last_names = ['احمدی', 'محمدی', 'رضایی', 'حسینی', 'کریمی', 'موسوی', 'نوری', 'صادقی', 'جعفری', 'اکبری',
                          'زمانی', 'کاظمی', 'شریفی', 'نظری', 'قاسمی', 'مرادی', 'یوسفی', 'باقری', 'طاهری', 'مهدوی']

            soldiers_created = 0
            today = timezone.now().date()
            dispatch_values = [choice[0] for choice in DispatchType.choices]

            for _ in range(50):
                # unique national_code
                while True:
                    national_code = str(random.randint(1000000000, 9999999999))
                    if not Soldier.objects.filter(national_code=national_code).exists():
                        break

                service_start = today - timedelta(days=random.randint(30, 600))
                birth_date = service_start - timedelta(days=random.randint(7000, 9000))

                if random.choice([True, False]):
                    arrival_date = service_start + timedelta(days=random.randint(0, 30))
                else:
                    arrival_date = None

                dispatch_type = random.choice(dispatch_values)

                soldier = Soldier(
                    first_name=random.choice(first_names),
                    last_name=random.choice(last_names),
                    national_code=national_code,
                    dossier_code=0,  # auto-assign positive number in save()
                    dispatch_type=dispatch_type,
                    service_start_date=service_start,
                    arrival_date=arrival_date,
                    birth_date=birth_date,
                    service_status=random.choice(active_statuses_for_seed),  # <<< only active
                    health_status=random.choice(health_statuses) if random.choice([True, False]) else None,
                    current_location=random.choice(locations) if random.choice([True, False]) else None,
                    phone_personal=f'09{random.randint(100000000, 999999999)}',
                    phone_parent=f'09{random.randint(100000000, 999999999)}',
                )

                # validate date logic from model
                soldier.full_clean()
                soldier.save()

                if expertises and random.choice([True, False]):
                    soldier.expertises.add(*random.sample(expertises, random.randint(1, 2)))

                if reduction_types and random.choice([True, False]):
                    reduction_type = random.choice(reduction_types)
                    SoldierServiceReduction.objects.create(
                        soldier=soldier,
                        reduction_type=reduction_type,
                        days=random.randint(5, reduction_type.default_days or 5),
                    )

                if random.choice([True, False]):
                    leave_start = today - timedelta(days=random.randint(0, 180))
                    if leave_start <= service_start:
                        leave_start = service_start + timedelta(days=random.randint(1, 10))
                    leave_end = leave_start + timedelta(days=random.randint(1, 5))

                    available_leave_types = [
                        LeaveType.MEDICAL,
                        LeaveType.REWARD,
                        LeaveType.EMERGENCY,
                    ]
                    if soldier.dispatch_type != DispatchType.ANSAR_PLAN:
                        available_leave_types.append(LeaveType.ENTITLEMENT)

                    leave_type = random.choice(available_leave_types)

                    try:
                        Leave.objects.create(
                            soldier=soldier,
                            leave_type=leave_type,
                            start_date=leave_start,
                            end_date=leave_end,
                            notes='مرخصی تستی',
                        )
                    except Exception as le:
                        logger.warning(f"Skipping invalid leave for soldier {soldier.pk}: {le}")

                soldiers_created += 1

            return JsonResponse({
                'success': True,
                'message': f'Successfully created {soldiers_created} soldiers with related data',
                'soldiers_created': soldiers_created,
            })

    except Exception as e:
        import traceback
        logger.error(f'Error seeding data: {e}\n{traceback.format_exc()}')
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@login_required
def panel_backup_restore(request):
    """Backup and restore page"""
    return render(request, 'dashboard/backup_restore.html')


@login_required
def export_database(request):
    """Export entire database to SQL dump file"""
    import subprocess
    import tempfile
    import os
    from django.conf import settings
    from django.utils import timezone

    try:
        db_config = settings.DATABASES['default']
        db_name = db_config['NAME']
        db_user = db_config['USER']
        db_password = db_config['PASSWORD']
        db_host = db_config['HOST']
        db_port = db_config['PORT']

        # Create temporary file for dump
        timestamp = timezone.now().strftime('%Y%m%d_%H%M%S')
        filename = f'soldier_db_backup_{timestamp}.sql'

        # Use temporary directory
        with tempfile.TemporaryDirectory() as tmpdir:
            dump_path = os.path.join(tmpdir, filename)

            # Set PGPASSWORD environment variable
            env = os.environ.copy()
            env['PGPASSWORD'] = db_password

            # Build pg_dump command
            # Using --no-password to prevent password prompt
            cmd = [
                'pg_dump',
                '--host', db_host,
                '--port', str(db_port),
                '--username', db_user,
                '--no-password',
                '--verbose',
                '--clean',
                '--if-exists',
                '--format=plain',
                '--file', dump_path,
                db_name
            ]

            # Execute pg_dump
            result = subprocess.run(
                cmd,
                env=env,
                capture_output=True,
                text=True,
                timeout=300  # 5 minute timeout
            )

            if result.returncode != 0:
                logger.error(f'pg_dump failed: {result.stderr}')
                messages.error(request, f'خطا در ایجاد پشتیبان: {result.stderr[:500]}')
                return redirect('panel_backup_restore')

            # Read the dump file
            if not os.path.exists(dump_path):
                messages.error(request, 'فایل پشتیبان ایجاد نشد.')
                return redirect('panel_backup_restore')

            with open(dump_path, 'rb') as f:
                response = HttpResponse(f.read(), content_type='application/sql')
                response['Content-Disposition'] = f'attachment; filename="{filename}"'

            logger.info(f'Database backup exported by user {request.user.username}')
            return response

    except subprocess.TimeoutExpired:
        logger.error('pg_dump timeout')
        messages.error(request, 'زمان ایجاد پشتیبان به پایان رسید. لطفاً دوباره تلاش کنید.')
        return redirect('panel_backup_restore')
    except FileNotFoundError:
        logger.error('pg_dump not found')
        messages.error(request, 'ابزار pg_dump یافت نشد. لطفاً PostgreSQL client tools را نصب کنید.')
        return redirect('panel_backup_restore')
    except Exception as e:
        logger.error(f'Error exporting database: {e}')
        import traceback
        logger.error(traceback.format_exc())
        messages.error(request, f'خطا در ایجاد پشتیبان: {str(e)}')
        return redirect('panel_backup_restore')


@login_required
def import_database(request):
    """Import database from SQL dump file"""
    import subprocess
    import tempfile
    import os
    from django.conf import settings
    from django.db import connection

    if request.method != 'POST':
        messages.error(request, 'فقط درخواست POST مجاز است.')
        return redirect('panel_backup_restore')

    if 'backup_file' not in request.FILES:
        messages.error(request, 'لطفاً فایل پشتیبان را انتخاب کنید.')
        return redirect('panel_backup_restore')

    backup_file = request.FILES['backup_file']

    # Validate file extension
    if not backup_file.name.endswith('.sql'):
        messages.error(request, 'فایل باید با پسوند .sql باشد.')
        return redirect('panel_backup_restore')

    try:
        db_config = settings.DATABASES['default']
        db_name = db_config['NAME']
        db_user = db_config['USER']
        db_password = db_config['PASSWORD']
        db_host = db_config['HOST']
        db_port = db_config['PORT']

        # Save uploaded file temporarily
        with tempfile.NamedTemporaryFile(mode='wb', suffix='.sql', delete=False) as tmp_file:
            for chunk in backup_file.chunks():
                tmp_file.write(chunk)
            tmp_file_path = tmp_file.name

        try:
            # Set PGPASSWORD environment variable
            env = os.environ.copy()
            env['PGPASSWORD'] = db_password

            # Build psql command
            cmd = [
                'psql',
                '--host', db_host,
                '--port', str(db_port),
                '--username', db_user,
                '--dbname', db_name,
                '--no-password',
                '--file', tmp_file_path,
                '--quiet'
            ]

            # Execute psql
            result = subprocess.run(
                cmd,
                env=env,
                capture_output=True,
                text=True,
                timeout=600  # 10 minute timeout
            )

            if result.returncode != 0:
                error_msg = result.stderr or result.stdout
                logger.error(f'psql restore failed: {error_msg}')
                messages.error(request, f'خطا در بازگردانی پشتیبان: {error_msg[:200]}')
                return redirect('panel_backup_restore')

            # Close all database connections to ensure clean state
            connection.close()

            logger.info(f'Database restored by user {request.user.username}')
            messages.success(request, 'پشتیبان با موفقیت بازگردانی شد. لطفاً صفحه را رفرش کنید.')

        finally:
            # Clean up temporary file
            try:
                os.remove(tmp_file_path)
            except:
                pass

        return redirect('panel_backup_restore')

    except subprocess.TimeoutExpired:
        logger.error('psql restore timeout')
        messages.error(request, 'زمان بازگردانی پشتیبان به پایان رسید. لطفاً دوباره تلاش کنید.')
        return redirect('panel_backup_restore')
    except FileNotFoundError:
        logger.error('psql not found')
        messages.error(request, 'ابزار psql یافت نشد. لطفاً PostgreSQL client tools را نصب کنید.')
        return redirect('panel_backup_restore')
    except Exception as e:
        logger.error(f'Error importing database: {e}')
        import traceback
        logger.error(traceback.format_exc())
        messages.error(request, f'خطا در بازگردانی پشتیبان: {str(e)}')
        return redirect('panel_backup_restore')
