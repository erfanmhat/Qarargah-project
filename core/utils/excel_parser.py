from openpyxl import load_workbook
import re
from datetime import date, timedelta


# --- helpers ---
def normalize(v):
    if v is None:
        return None
    s = str(v).strip()
    s = s.replace("ي", "ی").replace("ك", "ک")
    s = re.sub(r"\s+", " ", s)
    return s or None


def split_name_and_dossier_code(s):
    """
    ردیف 6: "نام نام‌خانوادگی - کدپرونده" یا "نام - کد"
    خروجی: (full_name, dossier_code)
    """
    s = normalize(s)
    if not s:
        return None, None
    parts = [p.strip() for p in s.split("-") if p.strip()]
    if len(parts) >= 2:
        try:
            int(parts[-1])
            full_name = parts[0]
            dossier_code = parts[-1]
        except:
            try:
                int(parts[0])
                full_name = parts[-1]
                dossier_code = parts[0]
            except:
                full_name = s
                dossier_code = 0
        return full_name, dossier_code
    # اگر جداکننده نبود: کل رشته را نام می‌گیریم
    return s, None


def jalali_to_yyyymmdd(y, m, d):
    """فعلاً فقط فرمت استاندارد خروجی می‌دهد. (تبدیل به میلادی را در مرحله ذخیره انجام بدهیم)"""
    return f"{int(y):04d}{int(m):02d}{int(d):02d}"


def parse_jalali_str(s):
    """
    B/C ممکن است مثل 14040918 یا 1404/09/18 باشد.
    خروجی: (y,m,d) یا None
    """
    s = normalize(s)
    if not s:
        return None
    s = s.replace("-", "/")
    if re.fullmatch(r"\d{8}", s):
        return int(s[:4]), int(s[4:6]), int(s[6:8])
    if re.fullmatch(r"\d{4}/\d{1,2}/\d{1,2}", s):
        y, m, d = s.split("/")
        return int(y), int(m), int(d)
    return None


def add_days_jalali_yyyymmdd(start_yyyymmdd, days):
    """
    چون جلالی است، جمع روز روی جلالی دقیق نیست مگر با jdatetime.
    برای اینکه خروجی درست باشد، پیشنهاد می‌کنم در پروژه jdatetime نصب شود.
    این تابع را طوری نوشته‌ام که اگر jdatetime موجود باشد دقیق حساب کند،
    وگرنه همان C را خالی می‌گذارد تا خطا/هشدار بدهیم.
    """
    try:
        import jdatetime
        y = int(start_yyyymmdd[:4])
        m = int(start_yyyymmdd[4:6])
        d = int(start_yyyymmdd[6:8])
        jd = jdatetime.date(y, m, d)
        # اگر مدت مرخصی 1 روز باشد، end همان start است => + (days-1)
        jd2 = jd + (days - 1)
        g = jd2.togregorian()  # صرفاً برای امنیت، ولی ما خروجی جلالی می‌خواهیم:
        # برگرداندن جلالی:
        return f"{jd2.year:04d}{jd2.month:02d}{jd2.day:02d}"
    except Exception:
        return None


def fix_date_to_jalali_yyyySmmSdd(yyyymmdd):
    try:
        yyyy = yyyymmdd[:4]
        mm = yyyymmdd[4:6]
        dd = yyyymmdd[6:8]
        return f"{yyyy}/{mm}/{dd}"
    except:
        return yyyymmdd


def parse_sheet_fixed(ws, sheet_name):
    """
    خواندن بر اساس مختصات ثابت:
    - ردیف 6 (Excel): B6 => نام - کد پرونده
    - ردیف 3 (Excel): G3,H3,I3 => روز/ماه/سال اعزام
    - مرخصی‌ها: از ردیف 8 به بعد: A,B,C,D
    """

    # --- Soldier header ---
    raw_header = ws["B6"].value  # ردیف 6، ستون B
    full_name, dossier_code = split_name_and_dossier_code(raw_header)

    # --- Dispatch date (G3,H3,I3) ---
    day = ws["G3"].value
    month = ws["H3"].value
    year = ws["I3"].value

    dispatch_date = None
    if day is not None and month is not None and year is not None:
        try:
            dispatch_date = fix_date_to_jalali_yyyySmmSdd(
                jalali_to_yyyymmdd(year, month, day)
            )

        except Exception:
            dispatch_date = None

    soldier = {
        "sheet_name": sheet_name,
        "full_name": full_name,
        "dossier_code": dossier_code,
        "dispatch_jalali": dispatch_date,
    }

    # --- Leaves from row 8 onward ---
    leaves = []
    row = 8
    while True:
        leave_type_fa = normalize(ws[f"A{row}"].value)
        start_raw = ws[f"B{row}"].value
        end_raw = ws[f"C{row}"].value
        duration_raw = ws[f"D{row}"].value
        print("=" * 80)
        print(leave_type_fa)
        print(start_raw)
        print(end_raw)
        print(duration_raw)
        print("#" * 80)

        # شرط پایان: اگر نوع مرخصی خالی شد و بقیه هم خالی، خروج
        if not leave_type_fa and start_raw is None and end_raw is None and duration_raw is None:
            break

        # اگر این ردیف واقعاً داده مرخصی نیست، ازش رد شو
        if not leave_type_fa:
            row += 1
            continue

        # start
        start_parts = parse_jalali_str(start_raw)
        start_yyyymmdd = None
        if start_parts:
            start_yyyymmdd = jalali_to_yyyymmdd(*start_parts)

        # duration
        duration_days = None
        try:
            if duration_raw is not None and str(duration_raw).strip() != "":
                duration_days = int(float(duration_raw))
        except Exception:
            duration_days = None

        # end (اگر خالی بود، محاسبه کن)
        end_yyyymmdd = None
        end_parts = parse_jalali_str(end_raw)
        if end_parts:
            end_yyyymmdd = jalali_to_yyyymmdd(*end_parts)
        else:
            if start_yyyymmdd and duration_days and duration_days > 0:
                end_yyyymmdd = add_days_jalali_yyyymmdd(start_yyyymmdd, duration_days)

        leaves.append({
            "leave_type_fa": leave_type_fa,
            "start_jalali": fix_date_to_jalali_yyyySmmSdd(start_yyyymmdd),
            "end_jalali": fix_date_to_jalali_yyyySmmSdd(end_yyyymmdd),
            "duration_days": duration_days,
            "row": row,  # برای گزارش خطا
        })

        row += 1

    return soldier, leaves, []


def parse_workbook(file_obj):
    wb = load_workbook(file_obj, data_only=True)
    parsed = []
    for name in wb.sheetnames:
        ws = wb[name]
        soldier, leaves, reductions = parse_sheet_fixed(ws, name)
        parsed.append({
            "sheet_name": name,
            "soldier": soldier,
            "leaves": leaves,
            "reductions": reductions,
        })
    return parsed
