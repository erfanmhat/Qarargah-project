from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from core.utils.excel_helpers import fmt_jdt, fmt_jdate


def _bytes_of_wb(wb: Workbook) -> bytes:
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _apply_rtl_sheet_defaults(ws, title: str):
    ws.title = title

    # RTL view (works in most Excel versions)
    ws.sheet_view.rightToLeft = True

    # Default fonts: Excel will use installed fonts; set a common Persian-friendly font name
    base_font = Font(name="IRANSans", size=11)  # if not installed, Excel will fallback
    header_font = Font(name="IRANSans", size=12, bold=True)

    # Common styles
    thin = Side(style="thin", color="666666")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_r = Alignment(horizontal="right", vertical="center", wrap_text=True)
    align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_fill = PatternFill("solid", fgColor="E7EEF8")

    return {
        "base_font": base_font,
        "header_font": header_font,
        "border": border,
        "align_r": align_r,
        "align_c": align_c,
        "header_fill": header_fill,
    }


def _set_col_widths(ws, widths: dict):
    for col_idx, w in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w


def _rank_display(soldier) -> str:
    """نمایش درجه نظامی با اولویت نام از جدول رفرنس."""
    if getattr(soldier, 'military_rank_ref', None):
        try:
            ref = soldier.military_rank_ref
            return f"{getattr(ref, 'title', '-') or '-'} ({getattr(ref, 'code', '-') or '-'})"
        except Exception:
            pass
    code = getattr(soldier, 'military_rank', None)
    return str(code) if code is not None else "-"


def _extra_service_display(soldier) -> str:
    """نمایش اضافه خدمت (روز + نوع)"""
    days = getattr(soldier, 'extra_service_days', None) or 0
    t = getattr(soldier, 'extra_service_type', None)
    type_title = getattr(t, 'title', None) if t else None

    if not days and not type_title:
        return "-"

    if type_title and days:
        return f"{days} روز - {type_title}"
    if days:
        return f"{days} روز"
    return type_title or "-"


def build_soldier_profile_excel(request, soldier, balances, now) -> bytes:
    """Kept name for backward compatibility with existing views.
    Returns XLSX bytes.
    """
    wb = Workbook()
    ws = wb.active
    st = _apply_rtl_sheet_defaults(ws, "پروفایل")

    # Title
    ws.merge_cells("A1:D1")
    c = ws["A1"]
    c.value = "پروفایل سرباز"
    c.font = Font(name="IRANSans", size=14, bold=True)
    c.alignment = st["align_c"]

    ws.merge_cells("A2:D2")
    c = ws["A2"]
    c.value = f"تاریخ صدور: {fmt_jdt(now)}"
    c.font = st["base_font"]
    c.alignment = st["align_c"]

    # Info rows (✅ آپدیت با فیلدهای جدید)
    fields = [
        ("کد سیستم", getattr(soldier, "id", "-") or "-"),
        ("نام و نام خانوادگی", getattr(soldier, "full_name", "-") or "-"),
        ("نام پدر", getattr(soldier, "father_name", "-") or "-"),
        ("کد ملی", getattr(soldier, "national_code", "-") or "-"),
        ("کد پرونده", getattr(soldier, "dossier_code", "-") or "-"),
        ("نوع اعزامی", getattr(soldier, "get_dispatch_type_display", lambda: "-")() or "-"),
        ("درجه نظامی", _rank_display(soldier)),
        ("وضعیت تاهل", getattr(soldier, "get_marital_status_display", lambda: "-")() if getattr(soldier, 'marital_status', None) else "-"),
        ("تعداد فرزندان", getattr(soldier, "children_count", "-") or "-"),
        ("تاریخ تولد", fmt_jdate(getattr(soldier, "birth_date", None))),
        ("محل تولد", getattr(soldier, "birth_place", "-") or "-"),
        ("محل صدور", getattr(soldier, "issue_place", "-") or "-"),
        ("قد (cm)", getattr(soldier, "height_cm", "-") or "-"),
        ("وزن (kg)", getattr(soldier, "weight_kg", "-") or "-"),
        ("رنگ چشم", getattr(soldier, "eye_color", "-") or "-"),
        ("گروه خونی", getattr(soldier, "get_blood_group_display", lambda: "-")() if getattr(soldier, 'blood_group', None) else "-"),
        ("بیماری زمینه‌ای", getattr(soldier, "underlying_disease", "-") or "-"),
        ("حساسیت دارویی", getattr(soldier, "drug_allergy", "-") or "-"),
        ("داروهای مصرفی", getattr(soldier, "used_medications", "-") or "-"),
        ("تحصیلات", getattr(soldier, "education_level", "-") or "-"),
        ("رشته تحصیلی", getattr(soldier, "education_field", "-") or "-"),
        ("گواهینامه رانندگی", "دارد" if getattr(soldier, 'has_driving_license', None) is True else ("ندارد" if getattr(soldier, 'has_driving_license', None) is False else "-")),
        ("نوع گواهینامه", getattr(soldier, "get_driving_license_type_display", lambda: "-")() if getattr(soldier, 'driving_license_type', None) else "-"),
        ("آدرس منزل پدری", getattr(soldier, "father_home_address", "-") or "-"),
        ("آدرس منزل شخصی", getattr(soldier, "personal_home_address", "-") or "-"),
        ("کد پستی", getattr(soldier, "postal_code", "-") or "-"),
        ("شماره تلفن منزل پدری", getattr(soldier, "father_home_phone", "-") or "-"),
        ("شماره تلفن منزل شخصی", getattr(soldier, "personal_home_phone", "-") or "-"),
        ("شماره تلفن همراه شخصی", getattr(soldier, "personal_mobile_phone", "-") or "-"),
        ("شماره تلفن نظامی", getattr(soldier, "military_mobile_phone", "-") or "-"),
        ("شماره همراه", getattr(soldier, "phone_personal", "-") or "-"),
        ("شماره والدین/همسر", getattr(soldier, "phone_parent", "-") or "-"),
        ("تاریخ شروع خدمت", fmt_jdate(getattr(soldier, "service_start_date", None))),
        ("تاریخ ورود به ناحیه", fmt_jdate(getattr(soldier, "arrival_date", None))),
        ("تاریخ پایان خدمت", fmt_jdate(getattr(soldier, "service_end_date", lambda: None)())),
        ("اضافه خدمت", _extra_service_display(soldier)),
        ("کد پاسداری", getattr(soldier, "guard_code", "-") or "-"),
        ("فرمانده مستقیم", f"{getattr(soldier, 'direct_commander_name', '-') or '-'} - {getattr(soldier, 'direct_commander_phone', '-') or '-'}"),
        ("دوره‌های گذرانده شده", getattr(soldier, "completed_courses", "-") or "-"),
    ]

    start_row = 4
    ws["A3"].value = "اطلاعات"
    ws["A3"].font = st["header_font"]
    ws["A3"].alignment = st["align_r"]

    r = start_row
    for label, value in fields:
        ws[f"C{r}"].value = str(label)
        ws[f"C{r}"].font = st["header_font"]
        ws[f"C{r}"].alignment = st["align_r"]
        ws[f"C{r}"].border = st["border"]

        ws[f"A{r}"].value = "" if value is None else str(value)
        ws[f"A{r}"].font = st["base_font"]
        ws[f"A{r}"].alignment = st["align_r"]
        ws[f"A{r}"].border = st["border"]

        ws.merge_cells(f"A{r}:B{r}")
        ws.merge_cells(f"C{r}:D{r}")
        r += 1

    # Balances table
    r += 1
    ws[f"A{r}"].value = "مانده مرخصی‌ها"
    ws[f"A{r}"].font = st["header_font"]
    ws[f"A{r}"].alignment = st["align_r"]
    ws.merge_cells(f"A{r}:D{r}")
    r += 1

    headers = ["نوع", "مجاز", "مصرف شده", "مانده"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=r, column=i, value=h)
        cell.font = st["header_font"]
        cell.fill = st["header_fill"]
        cell.alignment = st["align_c"]
        cell.border = st["border"]
    r += 1

    type_title = {"entitlement": "استحقاقی", "medical": "استعلاجی", "reward": "تشویقی"}
    for k in ("entitlement", "medical", "reward"):
        item = (balances or {}).get(k) or {}
        row = [
            type_title.get(k, k),
            item.get("allowed", 0),
            item.get("used", 0),
            item.get("remaining", 0),
        ]
        for col, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = st["base_font"]
            cell.alignment = st["align_c"] if col != 1 else st["align_r"]
            cell.border = st["border"]
        r += 1

    _set_col_widths(ws, {1: 26, 2: 14, 3: 18, 4: 18})
    ws.freeze_panes = "A4"

    return _bytes_of_wb(wb)


def build_soldier_leaves_excel(request, soldier, leaves, balances, now) -> bytes:
    """Kept name for backward compatibility with existing views.
    Returns XLSX bytes.
    """
    wb = Workbook()
    ws = wb.active
    st = _apply_rtl_sheet_defaults(ws, "مرخصی‌ها")

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "گزارش مرخصی‌های سرباز"
    c.font = Font(name="IRANSans", size=14, bold=True)
    c.alignment = st["align_c"]

    ws.merge_cells("A2:F2")
    c = ws["A2"]
    c.value = f"سرباز: {getattr(soldier, 'full_name', '-') or '-'}     تاریخ: {fmt_jdt(now)}"
    c.font = st["base_font"]
    c.alignment = st["align_c"]

    # Balances section
    row = 4
    ws.merge_cells(f"A{row}:F{row}")
    ws[f"A{row}"].value = "خلاصه مانده مرخصی‌ها"
    ws[f"A{row}"].font = st["header_font"]
    ws[f"A{row}"].alignment = st["align_r"]
    row += 1

    bal_headers = ["نوع", "مجاز", "مصرف", "مانده"]
    for i, h in enumerate(bal_headers, start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = st["header_font"]
        cell.fill = st["header_fill"]
        cell.alignment = st["align_c"]
        cell.border = st["border"]
    # merge E,F for balance header row aesthetics
    ws.merge_cells(f"E{row}:F{row}")
    row += 1

    type_title = {"entitlement": "استحقاقی", "medical": "استعلاجی", "reward": "تشویقی"}
    for k in ("entitlement", "medical", "reward"):
        item = (balances or {}).get(k) or {}
        vals = [
            type_title.get(k, k),
            item.get("allowed", 0),
            item.get("used", 0),
            item.get("remaining", 0),
        ]
        for col, val in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = st["base_font"]
            cell.alignment = st["align_c"] if col != 1 else st["align_r"]
            cell.border = st["border"]
        ws.merge_cells(f"E{row}:F{row}")
        row += 1

    # Leaves table
    row += 2
    ws.merge_cells(f"A{row}:F{row}")
    ws[f"A{row}"].value = "لیست مرخصی‌ها"
    ws[f"A{row}"].font = st["header_font"]
    ws[f"A{row}"].alignment = st["align_r"]
    row += 1

    headers = ["ردیف", "نوع", "از تاریخ", "تا تاریخ", "روز", "توضیحات"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = st["header_font"]
        cell.fill = st["header_fill"]
        cell.alignment = st["align_c"]
        cell.border = st["border"]
    row += 1

    for idx, lv in enumerate(leaves, start=1):
        leave_type_label = getattr(lv, "get_leave_type_display", lambda: getattr(lv, "leave_type", "-"))()
        vals = [
            idx,
            leave_type_label,
            fmt_jdate(getattr(lv, "start_date", None)),
            fmt_jdate(getattr(lv, "end_date", None)),
            getattr(lv, "days", "-") or "-",
            (getattr(lv, "notes", "") or "-").strip() or "-",
        ]
        for col, val in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = st["base_font"]
            cell.alignment = st["align_c"] if col != 6 else st["align_r"]
            cell.border = st["border"]
        row += 1

    _set_col_widths(ws, {1: 7, 2: 16, 3: 14, 4: 14, 5: 8, 6: 45})
    ws.freeze_panes = "A6"

    return _bytes_of_wb(wb)
